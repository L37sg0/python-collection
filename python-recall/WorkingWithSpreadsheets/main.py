import openpyxl

inv_file = openpyxl.load_workbook("inventory.xlsx")
product_list = inv_file["Sheet1"]

# List each company with respective product count
products_per_supplier = {}
# List each company with respective total inventory value
total_value_per_supplier = {}
# List products with inventory less than 10
products_under_10_inv = {}

for product_row in range(product_list.min_row + 1, product_list.max_row + 1):
    supplier_name = product_list.cell(product_row, 4).value
    inventory = product_list.cell(product_row, 2).value
    price = product_list.cell(product_row, 3).value
    product_number = product_list.cell(product_row,1).value
    inventory_price_total = product_list.cell(product_row,5)

    # calculate number of products per supplier
    if supplier_name in products_per_supplier:
        products_per_supplier[supplier_name] += 1
    else:
        products_per_supplier[supplier_name] = 1

    # calculate total value of inventory per supplier
    if supplier_name in total_value_per_supplier:
        # total_value_per_supplier.get(supplier_name)
        total_value_per_supplier[supplier_name] += inventory * price
    else:
        total_value_per_supplier[supplier_name] = inventory * price

    # logic products with inventory less than 10
    if inventory < 10:
        products_under_10_inv[int(product_number)] = int(inventory)

    # logic for adding a value for total inventory price for each product
    inventory_price_total.value = inventory * price
    product_list.cell(product_list.min_row,5).value = "Total inventory"

print(products_per_supplier)
print(total_value_per_supplier)
print(products_under_10_inv)
# Save the new file
inv_file.save("inventory_with_total_value.xlsx")

